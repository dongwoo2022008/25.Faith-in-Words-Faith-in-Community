import pandas as pd, re, numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
df=pd.read_pickle("lc_final.pkl"); H=pd.read_pickle("lc_hits_strict.pkl")
L=df[df.relig_strict].copy()
txt=(L['title'].fillna('')+' || '+L['desc_c'])
allrx=r"\bgod('s)?\b|\bjesus\b|\bchrist\b|\bchristians?\b|\bchurch(es)?\b|\bpray(s|ed|ing|er|ers)?\b|\bbless(ed|ing|ings)?\b|\bfaith(ful)?\b|\bthe lord\b|\bgood lord\b|\blord willing\b|\bbible\b|\bscripture\b|\btith(e|es|ing|er)\b|\bministr(y|ies)\b|\bmissionar(y|ies)\b|\bmission trip\b|\bpastor\b|\bcongregation\b|\bclergy\b|\brabbi\b|\bamen\b|\breligio(n|us)\b|\bcatholic\b|\bbaptist\b|\bmethodist\b|\blutheran\b|\bpresbyterian\b|\bmormon\b|\bevangelical\b|\bpentecostal\b|\bmuslim\b|\bislam(ic)?\b|\bmosque\b|\ballah\b|\bjewish\b|\bsynagogue\b|\bjudaism\b|\bbuddhis(t|m)\b|\bhindu\b|\blord\b"
def matched(t): return "; ".join(sorted(set(m.group(0).lower() for m in re.finditer(allrx,t,re.I))))
def ctx(t):
    m=re.search(allrx,t,re.I)
    if m is None: return t[:200]
    a=max(0,m.start()-100); b=min(len(t),m.end()+100); return "…"+t[a:b]+"…"
# rule-based first-pass type: (a) formulaic closing, (b) identity/value claim, (c) practice/expense
def typ(t):
    tl=t.lower()
    c=bool(re.search(r"\btith|\bchurch offering|\bmission trip|\bmissionar|\bpastor\b|\bministr|\bcongregation|\bclergy|\brabbi|\bchristian school|\bmy church\b|\bour church\b", tl))
    b=bool(re.search(r"\bas a christian\b|\bchristian (values|principles|family|man|woman)\b|\bi am a christian\b|\bi'?m a christian\b|\bbeliev(e|er)\b.*\b(god|christ|jesus)\b|\b(god|jesus|christ|the lord)\b.*\b(provide|will|has|led|steward|blessed us|grace)\b|\bmy faith\b|\bfaith in god\b|\bstrong faith\b", tl))
    a=bool(re.search(r"\bgod bless\b|\bgod bless you\b|\bthank god\b|\bthanks god\b|\bthank you god\b|\bgod help me\b|\bblessed (to|with)\b|\ba blessing\b|\bblessing\b|\bpray(ing)? (that|for|to|my|the)\b|\bprayer answered\b|\banswer to (a|my) prayer", tl))
    labs=[k for k,v in [('c 실천/지출',c),('b 정체성/가치',b),('a 관용구',a)] if v]
    return labs[0] if labs else 'a 관용구(추정)'
def frame(t):
    tl=t.lower(); f=[]
    if re.search(r"\bhonest|\bpromise|\bintegrity|\bresponsib|\btrustworthy|\bcommit|\bmy word\b|\bhonou?r\b|\bnever (been )?late|\bnever missed",tl): f.append('도덕')
    if re.search(r"\bstruggl|\bhard time|\blost my job|\blaid off|\bunemploy|\bmedical|\bdivorce|\bsick\b|\bhospital|\bemergency|\bbehind on|\bcan'?t afford|\bdesperate|\bsurgery|\billness|\bdisab",tl): f.append('곤경')
    return "/".join(f) if f else ''
term=L.loan_status.isin(['Fully Paid','Charged Off','Default'])
lc=pd.DataFrame({'loan_id':L['id'].astype(int),'year':L.year.astype('Int64'),'grade':L.grade,'loan_status':L.loan_status,'default(1/0/NA)':np.where(term, L.loan_status.isin(['Charged Off','Default']).astype(int), None),
  'desc_words':L.desc_len,'matched_terms':txt.map(matched),'context':txt.map(ctx),'auto_type_1st':txt.map(typ),'auto_frame_1st':txt.map(frame),'full_text':txt})
lc=lc.sort_values(['year','loan_id']).reset_index(drop=True)
# popfunding
pf=pd.read_pickle("pf.pkl"); P=pf[pf.core].copy()
KW=r'하나님께기도|기도해봅|하나님|하느님|예수|그리스도|주님|교회|(?<![가-힣])기도(?!원)|새벽기도|기도드리|신앙|성경|기독교|크리스천|목사|목회|전도사|선교|신학교|신학대|신학원|천주교|성당|수녀|신부님|(?<!선)불교(?!정|통)|스님|사찰|법당|부처님|종교|십일조|헌금|찬양|간증|성령|십자가'
def kmatched(t): return "; ".join(sorted(set(m.group(0) for m in re.finditer(KW,t))))
def kctx(t):
    m=re.search(KW,t)
    if m is None: return t[:160]
    a=max(0,m.start()-80); b=min(len(t),m.end()+80); return "…"+t[a:b].replace("\n"," ")+"…"
def ktyp(t):
    c=bool(re.search(r'십일조|헌금|목사|목회|전도사|선교|신학|교회\s?(직원|사무|간사|사역)',t)); b=bool(re.search(r'신앙|하나님(을|이|께서|의|만|뿐)|주님(을|께서|의)|믿는|기독교인|크리스천|신앙인|성경',t)); a=bool(re.search(r'기도(합니다|드립니다|하겠|하며|하고)|축복|감사드립니다',t))
    labs=[k for k,v in [('c 실천/지출',c),('b 정체성/가치',b),('a 관용구',a)] if v]; return labs[0] if labs else '검토필요'
def kframe(t):
    f=[]
    if re.search(r'성실|약속|책임|정직|신뢰|믿어주|반드시|절대|꼭 갚',t): f.append('도덕')
    if re.search(r'힘들|어려|병원|수술|실직|퇴사|폐업|사고|이혼|아프|암|장애|연체|신용불량|신불',t): f.append('곤경')
    return "/".join(f) if f else ''
pfd=pd.DataFrame({'listing_id':P['index'],'year':P.year,'status':P.status,'funded(1/0)':P.status.isin(['정상종료','연체중','상환중','손실(상환완료)','상환지연중']).astype(int),'bad_outcome(1/0/NA)':np.where(P.status.isin(['정상종료','연체중','상환중','손실(상환완료)','상환지연중']), P.status.isin(['연체중','손실(상환완료)','상환지연중']).astype(int), None),
  'text_chars':P.txt.str.len(),'matched_terms':P.txt.map(kmatched),'context':P.txt.map(kctx),'auto_type_1st':P.txt.map(ktyp),'auto_frame_1st':P.txt.map(kframe),'full_text':P.txt.str.replace("\n"," ")}).sort_values(['year','listing_id']).reset_index(drop=True)
print(len(lc),len(pfd)); print(lc.auto_type_1st.value_counts()); print(pfd.auto_type_1st.value_counts())
wb=Workbook(); hdr=Font(bold=True,color="FFFFFF",name="Arial"); fill=PatternFill("solid",fgColor="1F3864"); yellow=PatternFill("solid",fgColor="FFFF00"); base=Font(name="Arial",size=10)
def sheet(ws, d, coder_cols):
    cols=list(d.columns)+coder_cols
    ws.append(cols)
    for c in range(1,len(cols)+1):
        ws.cell(1,c).font=hdr; ws.cell(1,c).fill=fill; ws.cell(1,c).alignment=Alignment(wrap_text=True,vertical='center')
    for r in d.itertuples(index=False):
        ws.append([None if (isinstance(v,float) and np.isnan(v)) else (v if not hasattr(v,'item') else v.item()) for v in r])
    n=len(d)+1
    for c in range(1,len(cols)+1):
        for rr in range(2,n+1): ws.cell(rr,c).font=base
    for c in range(len(d.columns)+1,len(cols)+1):
        for rr in range(2,n+1): ws.cell(rr,c).fill=yellow
    widths={'context':60,'full_text':80,'matched_terms':22}
    for i,cn in enumerate(cols,1): ws.column_dimensions[get_column_letter(i)].width=widths.get(cn,14 if cn not in coder_cols else 16)
    ws.freeze_panes='A2'; ws.auto_filter.ref=f"A1:{get_column_letter(len(cols))}{n}"
    dv1=DataValidation(type="list",formula1='"a 관용구,b 정체성/가치,c 실천/지출,x 비종교(오탐)"',allow_blank=True); dv2=DataValidation(type="list",formula1='"종교,도덕,곤경,종교/도덕,종교/곤경,도덕/곤경,종교/도덕/곤경,해당없음"',allow_blank=True)
    ws.add_data_validation(dv1); ws.add_data_validation(dv2)
    ci=len(d.columns)
    for k,dv in [(1,dv1),(3,dv1),(2,dv2),(4,dv2)]:
        col=get_column_letter(ci+k); dv.add(f"{col}2:{col}{n}")
coder=['coder1_type','coder1_frame','coder2_type','coder2_frame','notes']
ws=wb.active; ws.title="README"
readme=[["신앙 표현 수동 코딩 시트 (2026-09-05 생성)"],[""],["목적","LendingClub 820건(desc 보유 717건 + 제목에만 표현 103건)·Popfunding 590건의 신앙 표현을 두 코더가 독립 코딩하여 κ를 산출하고, 유형별 분석 변수를 확정한다."],
["시트","LC_820: LendingClub title+desc 신앙 표현 전건 / PF_590: Popfunding 대출사유+상환계획 신앙 표현 전건 / 요약: 자동 1차 라벨 분포(참고용)"],
["노란색 셀","코더가 입력하는 칸. 드롭다운으로 선택."],["type 코드","a 관용구 = 'God bless', '기도합니다' 등 맺음말·관용 표현 / b 정체성·가치 = 신앙인으로서의 자기 규정·가치 주장 / c 실천·지출 = 십일조·헌금·교회 사역·선교 등 실제 행동·지출 보고 / x 비종교 = 오탐(예: 'Lord & Taylor')"],
["frame 코드","종교 = 종교 내용만 / 도덕 = 정직·약속·책임 언어 동반 / 곤경 = 실직·질병·연체 등 곤경 서술 동반 (복수 가능)"],
["auto_* 열","정규식 기반 1차 자동 라벨. 참고용이며 코더는 이를 보지 않고 코딩하는 것이 원칙(필요 시 열 숨김)."],
["예시 행","LC_820 시트 2행: matched_terms='god bless', context에 'Thank you and God bless'가 보이면 type=a 관용구, frame=해당없음"],
["출처","LendingClub accepted_2007_to_2018Q4.csv.gz; Popfunding popfunding.raw.list.20200130.txt (H:\\내 드라이브\\논문.김동우\\data\\P2P대출)"]]
for r in readme: ws.append(r)
ws['A1'].font=Font(bold=True,size=13,name="Arial"); ws.column_dimensions['A'].width=14; ws.column_dimensions['B'].width=120
for rr in range(1,len(readme)+1):
    for c in (1,2):
        if rr>1: ws.cell(rr,c).font=base
        ws.cell(rr,c).alignment=Alignment(wrap_text=True,vertical='top')
sheet(wb.create_sheet("LC_820"), lc, coder); sheet(wb.create_sheet("PF_590"), pfd, coder)
s=wb.create_sheet("요약"); s.append(["시트","auto_type_1st","건수"]); 
for nm,d in [("LC_820",lc),("PF_590",pfd)]:
    for k,v in d.auto_type_1st.value_counts().items(): s.append([nm,k,int(v)])
s.append([]); s.append(["LC_820 합계","=COUNTA(LC_820!A2:A100000)"]); s.append(["PF_590 합계","=COUNTA(PF_590!A2:A100000)"])
for c in range(1,4): s.cell(1,c).font=hdr; s.cell(1,c).fill=fill
s.column_dimensions['A'].width=14; s.column_dimensions['B'].width=22
wb.save("/mnt/user-data/outputs/신앙표현_수동코딩시트.xlsx"); print("saved")
