import subprocess, re, os
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH
body=open("manuscript_v17.md").read()
open("manuscript_v17_full.md","w").write(body)
subprocess.run(["pandoc","-f","markdown-tex_math_dollars","manuscript_v17_full.md","-o","manuscript_v17.docx","--resource-path=.","-V","geometry:margin=2.5cm"],check=True)
doc=Document("manuscript_v17.docx")
st=doc.styles['Normal']; st.font.name="Times New Roman"; st.font.size=Pt(11)
st.element.rPr.rFonts.set(qn('w:eastAsia'),"Times New Roman")
for s in doc.sections:
    s.page_width=Cm(21.0); s.page_height=Cm(29.7); s.left_margin=s.right_margin=Cm(2.5); s.top_margin=s.bottom_margin=Cm(2.5)
usable=s.page_width-s.left_margin-s.right_margin
for t in doc.tables:
    t.style=doc.styles['Table Grid'] if 'Table Grid' in [x.name for x in doc.styles] else t.style
    ncol=len(t.columns)
    hdr=[c.text.strip() for c in t.rows[0].cells]
    tblPr=t._tbl.tblPr
    lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
    tw=OxmlElement('w:tblW'); tw.set(qn('w:w'),str(int(usable/635))); tw.set(qn('w:type'),'dxa')
    for e in tblPr.findall(qn('w:tblW')): tblPr.remove(e)
    tblPr.append(tw)
    first=0.21 if ncol>8 else (0.24 if ncol>6 else (0.30 if ncol>3 else 0.45))
    if ncol==4: first=0.37
    widths=[usable*first]+[usable*(1-first)/(ncol-1)]*(ncol-1)
    fsize=7.5 if ncol>4 else 8.5
    if ncol==10 and hdr[0]=='Group':
        widths=[usable*0.20,usable*0.10]+[usable*0.70/8]*8
    if ncol==7 and hdr[1].startswith('(1) All years'):
        fr=[0.22,0.11,0.12,0.12,0.15,0.14,0.14]; widths=[usable*f for f in fr]; fsize=9
    if ncol==2:
        widths=[usable*0.55,usable*0.45]
    if ncol==4 and any('affiliated / comparison' in h for h in hdr):
        fr=[0.44,0.12,0.12,0.32]; widths=[usable*f for f in fr]
    grid=t._tbl.tblGrid
    for gc in list(grid): grid.remove(gc)
    for w in widths:
        gc=OxmlElement('w:gridCol'); gc.set(qn('w:w'),str(int(w/635))); grid.append(gc)
    borders=OxmlElement('w:tblBorders')
    for side,val in [('top','single'),('bottom','single'),('left','nil'),('right','nil'),('insideH','nil'),('insideV','nil')]:
        b=OxmlElement(f'w:{side}'); b.set(qn('w:val'),val); b.set(qn('w:sz'),'8'); b.set(qn('w:space'),'0'); b.set(qn('w:color'),'000000'); borders.append(b)
    for e in tblPr.findall(qn('w:tblBorders')): tblPr.remove(e)
    tblPr.append(borders)
    for cell in t.rows[0].cells:
        tcPr=cell._tc.get_or_add_tcPr(); tcb=OxmlElement('w:tcBorders'); b=OxmlElement('w:bottom'); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'6'); b.set(qn('w:color'),'000000'); tcb.append(b); tcPr.append(tcb)
    for row in t.rows:
        for i,cell in enumerate(row.cells):
            cell.width=int(widths[i])
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0)
                p.paragraph_format.keep_with_next=True; p.paragraph_format.keep_together=True
                for r in p.runs: r.font.size=Pt(fsize)
for p in doc.paragraphs:
    if p.style.name.startswith('Heading') or p.style.name in ('Title','Subtitle','Author','Date'):
        for r in p.runs: r.font.name="Times New Roman"; r.font.color.rgb=RGBColor(0,0,0)
    txt=p.text.strip()
    if txt.startswith("Notes."):
        for r in p.runs: r.font.size=Pt(9.5)
    if txt.startswith(("Panel ","Table ","Appendix Table ")):
        p.paragraph_format.keep_with_next=True
    if txt.startswith("Pr(D"):
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.save("manuscript_v17.docx")
subprocess.run(["soffice","--headless","--convert-to","pdf","--outdir","render","manuscript_v17.docx"],check=True,capture_output=True)
print("built")
